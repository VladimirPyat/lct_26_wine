from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[3]
DEFAULT_TEMPLATE = ROOT / "data" / "ЛЦТ2026 Шаблон презентации.pptx"
DEFAULT_LOGO = ROOT / "src" / "web" / "static" / "img" / "znak.png"
DEFAULT_OUTPUT = ROOT / "agent_docs" / "drafts" / "presentation_workflow" / "slide-copy-prototype.pptx"


def clone_slide(prs: Presentation, source_index: int):
    source = prs.slides[source_index]
    target = prs.slides.add_slide(source.slide_layout)

    relationship_ids: dict[str, str] = {}
    for rel_id, rel in source.part.rels.items():
        if rel.reltype == RT.SLIDE_LAYOUT:
            relationship_ids[rel_id] = next(
                candidate.rId
                for candidate in target.part.rels.values()
                if candidate.reltype == RT.SLIDE_LAYOUT
            )
            continue

        if rel.is_external:
            new_rel_id = target.part.rels._add_relationship(
                rel.reltype, rel.target_ref, True
            )
        else:
            new_rel_id = target.part.rels._add_relationship(
                rel.reltype, rel.target_part, False
            )
        relationship_ids[rel_id] = new_rel_id

    relationship_attributes = {qn("r:id"), qn("r:embed"), qn("r:link")}

    def remap_relationships(element) -> None:
        for attribute in relationship_attributes:
            old_rel_id = element.get(attribute)
            if old_rel_id in relationship_ids:
                element.set(attribute, relationship_ids[old_rel_id])
        for child in element:
            remap_relationships(child)

    source_sp_tree = source.shapes._spTree
    target_sp_tree = target.shapes._spTree
    for child in list(target_sp_tree)[2:]:
        target_sp_tree.remove(child)
    for child in list(source_sp_tree)[2:]:
        copied_child = deepcopy(child)
        remap_relationships(copied_child)
        target_sp_tree.append(copied_child)

    source_c_sld = source._element.cSld
    target_c_sld = target._element.cSld
    for attribute, value in source_c_sld.attrib.items():
        target_c_sld.set(attribute, value)
    for child in list(target_c_sld):
        if child is not target_sp_tree:
            target_c_sld.remove(child)
    for child in source_c_sld:
        if child is source_sp_tree:
            continue
        copied_child = deepcopy(child)
        remap_relationships(copied_child)
        target_c_sld.insert(target_c_sld.index(target_sp_tree), copied_child)

    return target


def replace_text(slide, old_text: str, new_text: str) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame or old_text not in shape.text:
            continue

        for paragraph in shape.text_frame.paragraphs:
            paragraph_text = "".join(run.text for run in paragraph.runs)
            if old_text not in paragraph_text:
                continue
            matching_run = next(
                (run for run in paragraph.runs if old_text in run.text), None
            )
            if matching_run is not None:
                matching_run.text = matching_run.text.replace(old_text, new_text)
            else:
                paragraph.runs[0].text = paragraph_text.replace(old_text, new_text)
                for run in paragraph.runs[1:]:
                    run.text = ""
            return

    raise ValueError(f"Text not found on the copied slide: {old_text!r}")


def keep_only_slide(prs: Presentation, slide) -> None:
    slide_ids = prs.slides._sldIdLst
    keep_rel_id = slide_ids[-1].rId
    for slide_id in list(slide_ids):
        if slide_id.rId == keep_rel_id:
            continue
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def set_page_number(slide, page_number: str) -> None:
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == 4:
            shape.text = page_number
            return
    raise ValueError("The copied slide has no page-number placeholder.")


def build_prototype(template: Path, logo: Path, output: Path) -> None:
    prs = Presentation(template)
    if len(prs.slides) < 4:
        raise ValueError("The template must contain at least four slides.")

    slide = clone_slide(prs, source_index=3)
    replace_text(slide, "Подробное описание решения", "Проверка копирования")
    replace_text(slide, "Маркетинговая часть решения", "Текст изменён")
    replace_text(slide, "Для продуктовых решений", "Скопировано из шаблона")
    set_page_number(slide, "0")

    slide.shapes.add_picture(
        str(logo), Inches(12.35), Inches(6.83), width=Inches(0.42)
    )
    keep_only_slide(prs, slide)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)

    check = Presentation(output)
    if len(check.slides) != 1:
        raise RuntimeError("Prototype output should contain exactly one slide.")
    visible_text = "\n".join(shape.text for shape in check.slides[0].shapes if shape.has_text_frame)
    for expected in ("Проверка копирования", "Текст изменён", "Скопировано из шаблона"):
        if expected not in visible_text:
            raise RuntimeError(f"Expected text is missing from output: {expected}")
    page_number = next(
        (shape for shape in check.slides[0].placeholders if shape.placeholder_format.idx == 4),
        None,
    )
    if page_number is None or page_number.text.strip() != "0":
        raise RuntimeError("The page-number placeholder was not changed to 0.")
    if not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in check.slides[0].shapes):
        raise RuntimeError("The inserted logo is missing from output.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Copy and edit a template slide as a PPTX proof of concept.")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--logo", type=Path, default=DEFAULT_LOGO)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    build_prototype(args.template, args.logo, args.output)
    print(f"Created {args.output}")


if __name__ == "__main__":
    main()